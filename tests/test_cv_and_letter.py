import re
import shutil
from pathlib import Path

import pytest
from docx import Document

from mp import cv as cvmod
from mp import letter as lettermod
from mp.models import CvEdit, Letter
from mp.pdf import docx_to_pdf, soffice_binary
from tests.conftest import FakeClaude


def test_select_profile_tree():
    assert cvmod.select_profile("Implementation of SimCorp Dimension front-to-back") == "AssetManagement"
    assert cvmod.select_profile("IFRS 17 and Solvency II reporting, QRT") == "IFRS17SolvencyII"
    assert cvmod.select_profile("IFRS 9 ECL on a SimCorp platform") == "AssetManagement"  # AM d'abord
    assert cvmod.select_profile("Finance transformation PMO") == "FinanceTransformation"
    assert cvmod.select_profile("Chef de projet", hint="IFRS17SolvencyII") == "IFRS17SolvencyII"
    assert cvmod.select_profile("Chef de projet", hint="n'importe quoi") == "FinanceTransformation"


def test_select_profile_whole_words_only():
    # Faux positifs de l'ancien MATT (sous-chaînes) : ABOR dans « collaboration », OMS dans « telecoms »,
    # ECL dans « declaration », PAA dans « paas ».
    ft = "FinanceTransformation"
    assert cvmod.select_profile("Transformation finance, forte collaboration avec les métiers") == ft
    assert cvmod.select_profile("Program manager, elaborate the roadmap for a telecoms client") == ft
    assert cvmod.select_profile("Tax declaration process redesign on a PaaS platform") == ft
    assert cvmod.select_profile("Implémentation d'un OMS pour une société de gestion") == "AssetManagement"
    assert cvmod.select_profile("ABOR / IBOR reconciliation") == "AssetManagement"
    assert cvmod.select_profile("Expected credit loss (ECL) model") == "IFRS17SolvencyII"
    assert cvmod.select_profile("Reporting prudentiel et provisions techniques") == "IFRS17SolvencyII"


def test_select_profile_hint_never_overrides_the_tree_when_the_posting_is_known():
    # 3 octobre : PMO Digital Finance KPMG (aucun mot-clé AM / IFRS) marqué « AssetManagement » → CV AM. Le CV 1 s'impose.
    kpmg = "Senior Manager – Digital Finance - PMO. Piloter le portefeuille de projets, gouvernance, budget, KPI."
    assert cvmod.select_profile(kpmg, hint="AssetManagement", has_jd=True) == "FinanceTransformation"
    assert cvmod.select_profile("Senior Manager PMO", hint="AssetManagement") == "AssetManagement"   # titre seul


def test_select_profile_hint_arbitrates_when_both_families_match():
    both = "IFRS 17 programme on a SimCorp platform"
    assert cvmod.select_profile(both) == "AssetManagement"                       # ordre du skill
    assert cvmod.select_profile(both, hint="IFRS17SolvencyII") == "IFRS17SolvencyII"
    assert cvmod.select_profile("SimCorp migration", hint="IFRS17SolvencyII") == "AssetManagement"


def _make_cv(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("XAVIER ROBITAILLE")
    p = doc.add_paragraph()
    p.add_run("Finance Transformation  |  ")
    r = p.add_run("PMO")
    r.bold = True
    p.add_run("  |  IFRS 9")
    doc.add_paragraph("Led end-to-end delivery of a Clearwater Analytics SaaS platform for a French reinsurer.")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).paragraphs[0].add_run("Core expertise: IFRS 9 / IFRS 17 implementation.")
    doc.save(str(path))
    return path


def test_apply_edits_single_run_multi_run_table_and_unknown(tmp_path):
    base = _make_cv(tmp_path / "base.docx")
    out = tmp_path / "out.docx"
    edits = [
        CvEdit(old="Clearwater Analytics SaaS platform", new="Clearwater Analytics SaaS investment platform"),
        CvEdit(old="PMO  |  IFRS 9", new="PMO  |  IFRS 9  |  IFRS 17"),          # chevauche 2 runs
        CvEdit(old="IFRS 9 / IFRS 17 implementation", new="IFRS 9 / IFRS 17 / Solvency II implementation"),
        CvEdit(old="texte qui n'existe pas", new="x"),
    ]
    applied, skipped = cvmod.apply_edits(base, out, edits)
    assert (applied, skipped) == (3, 1)
    text = cvmod.docx_text(out)
    assert "SaaS investment platform" in text
    assert "PMO  |  IFRS 9  |  IFRS 17" in text
    assert "Solvency II implementation" in text
    # la mise en forme du run « PMO » (gras) est conservée
    doc = Document(str(out))
    assert any(r.bold for r in doc.paragraphs[1].runs)


def test_propose_edits_filters_unsafe(tmp_path):
    base = _make_cv(tmp_path / "base.docx")
    claude = FakeClaude([{"edits": [
        {"old": "PMO", "new": "PMO " + "x" * 200},                       # dérive de longueur
        {"old": "", "new": "y"},                                         # vide
        {"old": "Finance Transformation", "new": "Finance Transformation (actuarial qualification)"},  # actuaire
        {"old": "Finance Transformation", "new": "Finance & Insurance Transformation"},
    ], "gaps": ["Aladdin non pratiqué"]}])
    plan = cvmod.propose_edits(claude, cv_text=cvmod.docx_text(base), profile="FinanceTransformation", lang="EN",
                               title="PMO", employer="AXA", jd_text="", keywords=["PMO"], profile_md="p")
    assert [e.new for e in plan.edits] == ["Finance & Insurance Transformation"]
    assert plan.gaps == ["Aladdin non pratiqué"]


def test_fallback_prompt_carries_tailoring_rules(tmp_path):
    # Repli API : mêmes règles de surfaçage que la skill cv-tailoring (migration, apurement, VBA, SAP).
    base = _make_cv(tmp_path / "base.docx")
    claude = FakeClaude([{"edits": [], "gaps": []}])
    cvmod.propose_edits(claude, cv_text=cvmod.docx_text(base), profile="FinanceTransformation", lang="FR",
                        title="Chef de projet migration comptable", employer="AXA", jd_text="", keywords=[],
                        profile_md="p")
    system = "\n".join(claude.prompts[0]["system"])
    for term in ("bascule d'interfaces comptables", "apurement de suspens", "macros VBA", "comptabilité auxiliaire"):
        assert term in system


def test_base_cv_path_fallback(tmp_path):
    (tmp_path / "CV_XRO_EN_FinanceTransformation_v5.docx").write_bytes(b"x")
    p, lang = cvmod.base_cv_path(tmp_path, "FinanceTransformation", "FR")
    assert lang == "EN" and p.name.startswith("CV_XRO_EN")
    with pytest.raises(FileNotFoundError):
        cvmod.base_cv_path(tmp_path, "AssetManagement", "EN")


def test_output_names():
    assert cvmod.output_name("FR", "AssetManagement", "Groupe P&V", "123", "20261001") == \
        "CV_XRO_FR_AssetManagement_GroupePV_123_20261001"
    assert lettermod.output_name("EN", "Swiss Re", "9", "20261001") == "CL_XRO_EN_SwissRe_9_20261001"
    assert lettermod.output_name("FR", "", "", "20261001") == "LM_XRO_FR_offre_20261001"


def test_geo_clauses():
    uk = lettermod.geo_clauses("London, United Kingdom", "", "Programme Manager")
    assert any("IR35" in c for c in uk) and any("HYBRIDE" in c for c in uk)
    fr = lettermod.geo_clauses("Paris", "Poste en CDI au sein de la direction financière", "Head of Finance")
    assert any("CDI" in c for c in fr) and not any("IR35" in c for c in fr)
    assert lettermod.geo_clauses("Paris", "Mission freelance de 6 mois, TJM à négocier", "PMO") == []
    assert any("HYBRIDE" in c for c in lettermod.geo_clauses("Dublin", "", ""))


def test_write_letter_retries_when_too_short():
    short = {"lettre": "Trop court. " * 10, "objections": []}
    ok_text = " ".join(["mot"] * 230)
    claude = FakeClaude([short, {"lettre": ok_text, "objections": ["TJM — à cadrer"]}])
    letter = lettermod.write_letter(claude, lang="FR", title="PMO", employer="AXA", location="Paris", jd_text="",
                                    profile_md="p", writing_rules="règles")
    assert len(claude.prompts) == 2
    assert "hors fourchette" in claude.prompts[1]["user"]
    assert isinstance(letter, Letter) and letter.objections == ["TJM — à cadrer"]
    assert claude.prompts[0]["system"][3] == "règles"


def test_writing_rules_sent_in_full():
    # L'ancien MATT chargeait les règles d'écriture en entier ; la v3 les tronquait à 16 000 caractères.
    rules = "Règle. " * 5000                      # ≈ 35 000 caractères, plus long que les vrais fichiers
    claude = FakeClaude([{"lettre": " ".join(["mot"] * 230), "objections": []}])
    lettermod.write_letter(claude, lang="FR", title="PMO", employer="AXA", location="Paris", jd_text="",
                           profile_md="p", writing_rules=rules)
    assert claude.prompts[0]["system"][3] == rules
    real = lettermod.load_writing_rules(Path(__file__).parents[1] / "assets" / "writing_rules", "FR")
    assert len(real) > 16000                      # le fichier réel dépasse l'ancienne coupure


def test_letter_docx_and_pdf(tmp_path):
    text = ("L'enjeu auquel doit faire face la direction financière d'AXA est peu commun : un paragraphe.\n\n"
            "Chez CNP, j'ai automatisé la production financière projetée multi-actifs.\n\n"
            "Je serais heureux d'avoir l'opportunité d'en discuter avec vous.\n\nBien cordialement.")
    out = lettermod.letter_docx(text, tmp_path / "lm.docx", employer="AXA", title="PMO Finance", lang="FR")
    doc = Document(str(out))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert paras[0] == "Xavier Robitaille"
    assert "AXA" in paras and any(p.startswith("Objet : PMO Finance") for p in paras)
    assert paras[-1] == "Xavier Robitaille" and paras[-2] == "Bien cordialement."
    if not soffice_binary() or not shutil.which("pdfinfo"):
        pytest.skip("LibreOffice absent")
    pdf = docx_to_pdf(out, tmp_path)
    assert pdf.exists() and pdf.stat().st_size > 5000


def test_letter_signed_once_when_text_already_signed(tmp_path):
    # La skill cover-letter termine la lettre par « Xavier Robitaille » (dossier Alpha FMC du 3 octobre) ; le DOCX
    # ajoute déjà la signature : le nom ne doit apparaître qu'une fois en fin de lettre.
    text = "Alpha's team in Paris.\n\nI would be happy to discuss.\n\nBest regards,\n\nXavier Robitaille\n"
    out = lettermod.letter_docx(text, tmp_path / "cl.docx", employer="Alpha FMC", title="Senior Consultant", lang="EN")
    paras = [p.text for p in Document(str(out)).paragraphs if p.text.strip()]
    assert paras[-2:] == ["Best regards,", "Xavier Robitaille"]
    assert paras.count("Xavier Robitaille") == 2          # en-tête + signature, pas de doublon
    assert lettermod.strip_signature("Corps.\n\nBien cordialement.") == "Corps.\n\nBien cordialement."


def test_word_count():
    assert lettermod.word_count("Bien cordialement. L'enjeu est peu commun — vraiment.") == 7  # L'enjeu = 1 mot


def test_letter_prompts_forbid_self_opening_and_copied_example():
    # Dossier Alpha FMC du 3 octobre : accroche qui glissait vers Xavier, exemple Swiss Re recopié. Repli API aligné
    # sur la skill cover-letter v2.6.
    from mp.config import prompt
    en, fr = prompt("cover_en"), prompt("cover_fr")
    assert 'contains no "I", "my" or "me"' in en and "never copy its wording" in en
    assert "ni « je », ni « mon », ni « mes »" in fr and "jamais l'exemple Swiss Re" in fr


# ---------------------------------------------------------------------------
# CV de base v5 (3 octobre 2026) : format du CV ABOR du 27 septembre, faits durs, dates arbitrées par Xavier
# ---------------------------------------------------------------------------

CV_BASE_DIR = Path(__file__).resolve().parents[1] / "assets" / "cv_base"
FORBIDDEN_CV = re.compile(r"\bactuary\b|\bactuaire\b|qualification actuarielle|actuarial qualification|"
                          r"master (in|en) (sciences )?actuari|advised investment coo|special advisor|\bAPAC\b|"
                          r"north america|amérique du nord", re.I)


@pytest.mark.parametrize("lang, profile", [(lg, pr) for lg in ("EN", "FR") for pr in cvmod.PROFILES])
def test_base_cvs_v5_respect_hard_facts(lang, profile):
    text = cvmod.docx_text(CV_BASE_DIR / cvmod.CV_FILES[lang][profile])
    bad = FORBIDDEN_CV.search(text)
    assert not bad, bad and bad.group(0)                      # faits durs 1 et 5, cap du 6 septembre
    assert "chief transformation officer" in text.lower()
    assert ("all examinations passed" if lang == "EN" else "ensemble des épreuves") in text
    assert "CNP Assurances  ·  2013 – 2017" in text           # arbitrage du 3 octobre
    assert ("Nov 2024 – June 2026" if lang == "EN" else "nov. 2024 – juin 2026") in text


def test_edits_apply_on_the_v5_layout_and_keep_the_keyword_bar_format(tmp_path):
    base = CV_BASE_DIR / cvmod.CV_FILES["FR"]["AssetManagement"]
    out = tmp_path / "cv.docx"
    edits = [CvEdit(old="SAS  |  Power BI", new="SAS  |  Power BI  |  Aladdin"),
             CvEdit(old="d'hebdomadaire à quotidienne", new="d'hebdomadaire à quotidienne (multi-dépositaires)")]
    assert cvmod.apply_edits(base, out, edits) == (2, 0)
    kw = next(p for p in Document(str(out)).paragraphs if "Power BI  |  Aladdin" in p.text)
    assert all(r.font.size.pt == 2 for r in kw.runs if r.text)   # barre de mots-clés toujours invisible à l'œil
