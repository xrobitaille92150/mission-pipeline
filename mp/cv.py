"""Sélection du CV de base, retouches chirurgicales proposées par Claude, application python-docx."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document

from mp.claude import Claude
from mp.config import prompt
from mp.models import CV_EDITS_SCHEMA, CvEdit, CvEditPlan

log = logging.getLogger("mp.cv")

PROFILES = ("FinanceTransformation", "AssetManagement", "IFRS17SolvencyII")

CV_FILES = {
    "EN": {
        "FinanceTransformation": "CV_XRO_EN_FinanceTransformation_v5.docx",
        "AssetManagement": "CV_XRO_EN_AssetManagement_v5.docx",
        "IFRS17SolvencyII": "CV_XRO_EN_IFRS17_SolvencyII_v5.docx",
    },
    "FR": {
        "FinanceTransformation": "CV_XRO_FR_FinanceTransformation_v5.docx",
        "AssetManagement": "CV_XRO_FR_AssetManagement_v5.docx",
        "IFRS17SolvencyII": "CV_XRO_FR_IFRS17_SolvencyII_v5.docx",
    },
}

# Arbre de décision de cv_profiles.md (skill cv-tailoring). L'ordre compte : AM avant IFRS.
AM_KW = ["oms", "simcorp", "clearwater", "aladdin", "bloomberg aim", "murex", "linedata", "wealthsuite",
         "front-to-back", "front to back", "investment platform", "investment accounting", "investment reporting",
         "comptabilité des placements", "comptabilité investissement", "reporting investissement", "abor", "ibor",
         "custodian", "dépositaire", "securities migration", "migration titres", "fund accounting",
         "comptabilité opc", "middle office", "back office", "asset servicing"]
IFRS_KW = ["ifrs 17", "ifrs17", "ifrs 9", "ifrs9", "solvency ii", "solvency 2", "solvabilité ii", "solvabilité 2",
           "qrt", "orsa", "actuarial", "actuariat", "technical provisions", "provisions techniques", "dry-run",
           "dry run", "paa", "gmm", "ecl", "reserving", "provisionnement", "comptabilité technique",
           "reporting prudentiel", "regulatory reporting", "pillar 3", "pilier 3", "rsr", "sfcr"]


def _kw_regex(words: list[str]) -> re.Pattern:
    # Mots entiers seulement : « abor » (ABOR) ne doit pas matcher « collaboration », ni « ecl » (ECL)
    # « declaration », ni « oms » (OMS) « telecoms ». \w couvre les lettres accentuées.
    return re.compile(r"(?<!\w)(?:" + "|".join(re.escape(w) for w in words) + r")(?!\w)")


AM_RE, IFRS_RE = _kw_regex(AM_KW), _kw_regex(IFRS_KW)


def select_profile(text: str, hint: str | None = None, has_jd: bool = False) -> str:
    """Profil CV : règle déterministe du skill cv-tailoring (mots entiers, AM avant IFRS ; aucun mot-clé = CV 1).
    `hint` (avis de Claude au scoring) ne fait que trancher quand les deux familles de mots-clés sont présentes,
    ou quand on n'a que le titre (`has_jd` faux). Avec une fiche de poste, il ne contredit jamais l'arbre :
    le 3 octobre, quatre postes de transformation finance / PMO (KPMG, Accenture, CSC, Allianz Services)
    portaient « AssetManagement » et avaient reçu le CV Asset Management."""
    t = (text or "").lower()
    am, ifrs = bool(AM_RE.search(t)), bool(IFRS_RE.search(t))
    if am and ifrs and hint in ("AssetManagement", "IFRS17SolvencyII"):
        return hint
    if am:
        return "AssetManagement"
    if ifrs:
        return "IFRS17SolvencyII"
    if hint in PROFILES and not has_jd:
        return hint
    return "FinanceTransformation"


def base_cv_path(cv_dir: Path, profile: str, lang: str) -> tuple[Path, str]:
    """Chemin du CV de base et langue effective (repli EN si le FR manque)."""
    lang = "FR" if lang == "FR" else "EN"
    p = cv_dir / CV_FILES[lang][profile]
    if lang == "FR" and not p.exists():
        log.warning("CV FR absent pour %s — repli sur la version EN", profile)
        lang = "EN"
        p = cv_dir / CV_FILES["EN"][profile]
    if not p.exists():
        raise FileNotFoundError(f"CV de base introuvable : {p}")
    return p, lang


def _paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in _paragraphs(doc) if p.text.strip())


def _replace_in_paragraph(para, old: str, new: str) -> bool:
    full = "".join(r.text for r in para.runs)
    if old not in full:
        return False
    # 1. cas simple : un run contient tout le texte à remplacer → mise en forme intacte
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    # 2. le texte chevauche plusieurs runs → on fusionne le segment dans le premier run touché
    start = full.index(old)
    end = start + len(old)
    pos = 0
    first = last = None
    for i, r in enumerate(para.runs):
        r_start, r_end = pos, pos + len(r.text)
        if first is None and r_end > start:
            first = i
        if r_start < end:
            last = i
        pos = r_end
    if first is None or last is None:
        return False
    runs = para.runs
    seg_start = sum(len(r.text) for r in runs[:first])
    seg_text = "".join(r.text for r in runs[first:last + 1])
    rel = start - seg_start
    runs[first].text = seg_text[:rel] + new + seg_text[rel + len(old):]
    for r in runs[first + 1:last + 1]:
        r.text = ""
    return True


def apply_edits(base: Path, out: Path, edits: list[CvEdit]) -> tuple[int, int]:
    """Applique les remplacements ; retourne (appliqués, ignorés). Un `old` introuvable est ignoré."""
    doc = Document(str(base))
    applied = skipped = 0
    for e in edits:
        if not e.old or e.old == e.new:
            skipped += 1
            continue
        done = False
        for p in _paragraphs(doc):
            if _replace_in_paragraph(p, e.old, e.new):
                done = True
                break
        if done:
            applied += 1
        else:
            skipped += 1
            log.info("retouche ignorée (texte introuvable) : %r", e.old[:60])
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return applied, skipped


def propose_edits(claude: Claude, *, cv_text: str, profile: str, lang: str, title: str, employer: str,
                  jd_text: str, keywords: list[str], profile_md: str) -> CvEditPlan:
    system = [profile_md, prompt("cv_edits")]
    user = (f"CV DE BASE (profil {profile}, langue {'français' if lang == 'FR' else 'anglais'}) — "
            f"chaque paragraphe sur une ligne :\n\n{cv_text}\n\n"
            f"OFFRE :\nPoste : {title}\nEmployeur : {employer}\n"
            f"Mots-clés repérés au scoring : {', '.join(keywords) or '—'}\n"
            f"Description :\n{(jd_text or '(non disponible — adapter sur le titre seul, 2 retouches max)')[:7000]}\n\n"
            "Propose les retouches (JSON).")
    data = claude.json(system=system, user=user, schema=CV_EDITS_SCHEMA, effort="medium", max_tokens=8000)
    plan = CvEditPlan.model_validate(data)
    # garde-fous : pas de retouche vide, pas de dérive de longueur, pas plus de 6
    safe = []
    for e in plan.edits[:6]:
        if not e.old.strip() or not e.new.strip():
            continue
        if len(e.new) > 1.6 * len(e.old) + 40:
            continue
        if re.search(r"\bactuai?r(e|y|ial)\b", e.new, re.I) and not re.search(r"\bactuai?r", e.old, re.I):
            continue
        safe.append(e)
    plan.edits = safe
    return plan


def output_name(lang: str, profile: str, employer: str, job_id: str, day: str) -> str:
    safe = re.sub(r"[^\w\-]", "", (employer or "offre").replace(" ", ""))[:30] or "offre"
    jid = f"_{job_id}" if job_id else ""
    return f"CV_XRO_{lang}_{profile}_{safe}{jid}_{day}"
