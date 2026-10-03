"""Lettre de motivation (FR) / cover text (EN) : prompt, génération, contrôle de longueur, DOCX."""
from __future__ import annotations

import datetime as dt
import logging
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from mp.claude import Claude
from mp.config import prompt
from mp.models import LETTER_SCHEMA, Letter

log = logging.getLogger("mp.letter")

BRAND_NAVY = RGBColor(0x0B, 0x15, 0x30)
BRAND_GOLD = RGBColor(0xC7, 0x9A, 0x3B)
INK_SOFT = RGBColor(0x5A, 0x62, 0x78)
CONTACT_NAME = "Xavier Robitaille"
CONTACT_ROLE = "FINANCE & INSURANCE TRANSFORMATION"
CONTACT_EMAIL = "xro@xavier-robitaille.com"
CONTACT_PHONE = "+33 6 64 89 09 43"

WORDS = {"FR": (180, 280), "EN": (220, 350)}
WRITING_RULES_FILES = {"FR": "REGLES-ECRITURE-FR.md", "EN": "WRITING RULES.md"}

IR35_SENTENCE = "As a contractor who is a non-UK Tax resident, based overseas, IR35 does not apply to me."
HYBRID_CLAUSE = ("REMOTE/HYBRIDE : intègre naturellement que le travail hybride est une pratique courante pour "
                 "Xavier et qu'il maîtrise les outils et pratiques du travail à distance (référence concrète : "
                 "sa dernière mission chez Clearwater, pilotée pour un client luxembourgeois).")

UK = ["united kingdom", "royaume-uni", "angleterre", "england", "scotland", "écosse", "ecosse", "wales", "london",
      "londres", "manchester", "edinburgh", "glasgow", "birmingham", "leeds", "bristol", "liverpool", " uk"]
IE = ["ireland", "irlande", "dublin", "cork", "galway", "limerick"]
BENELUX = ["belgium", "belgique", "belgië", "brussels", "bruxelles", "antwerp", "anvers", "gent", "ghent", "netherlands",
           "pays-bas", "nederland", "amsterdam", "rotterdam", "hague", "haye", "utrecht", "luxembourg", "benelux"]
FRANCE = ["france", "paris", "lyon", "marseille", "lille", "toulouse", "bordeaux", "nantes", "nice", "strasbourg",
          "rennes", "montpellier", "suresnes", "défense", "defense", "neuilly", "courbevoie", "levallois", "ville de paris"]
FREELANCE_SIG = ["freelance", "mission", "indépendant", "independant", "consultant externe", "prestation", "prestataire",
                 "contractor", "contract role", "b2b", "daily rate", "tjm", "interim", "intérim", "sous-traitance",
                 "management de transition", "portage"]


def geo_clauses(location: str, jd_text: str, title: str) -> list[str]:
    loc = (location or "").lower().strip() or (jd_text or "")[:400].lower()
    text_all = f"{title or ''} {jd_text or ''}".lower()
    uk = any(k in f" {loc}" for k in UK)
    ie = any(k in loc for k in IE)
    benelux = any(k in loc for k in BENELUX)
    france = any(k in loc for k in FRANCE)
    freelance = any(k in text_all for k in FREELANCE_SIG)
    clauses = []
    if uk:
        clauses.append(HYBRID_CLAUSE)
        clauses.append('IR35 : inclus la phrase EXACTE ci-dessous, en anglais, telle quelle — ne la traduis pas, '
                       f'ne la reformule pas, ne la coupe pas : "{IR35_SENTENCE}"')
    elif ie or benelux:
        clauses.append(HYBRID_CLAUSE)
    if france and not freelance:
        clauses.append("OUVERTURE CDI : ce poste est un rôle interne en France. Utilise la formulation validée du "
                       "fait dur n°9 (statut indépendant, CDI envisageable), sans lourdeur.")
    return clauses


def word_count(text: str) -> int:
    return len(re.findall(r"[\w'’-]+", text or ""))


def load_writing_rules(rules_dir: Path, lang: str) -> str:
    p = rules_dir / WRITING_RULES_FILES["FR" if lang == "FR" else "EN"]
    if p.exists():
        return p.read_text(encoding="utf-8")
    log.warning("règles d'écriture absentes (%s) — lettre générée avec les extraits embarqués", p)
    return ""


def write_letter(claude: Claude, *, lang: str, title: str, employer: str, location: str, jd_text: str,
                 profile_md: str, writing_rules: str, gaps: list[str] | None = None) -> Letter:
    lang = "FR" if lang == "FR" else "EN"
    system = [prompt("cover_common"), prompt("cover_fr" if lang == "FR" else "cover_en"),
              profile_md, (writing_rules or "")]
    clauses = geo_clauses(location, jd_text, title)
    extra = ""
    if clauses:
        extra = ("\n\nPRÉCISIONS À INTÉGRER (obligatoire, de façon fluide, jamais en liste) :\n"
                 + "\n".join(f"- {c}" for c in clauses))
    if gaps:
        extra += ("\n\nÉCARTS IDENTIFIÉS AU STADE DU CV (à traiter dans `objections`, jamais dans la lettre) :\n"
                  + "\n".join(f"- {g}" for g in gaps))
    user = (f"OFFRE\nEmployeur : {employer}\nPoste : {title}\nLieu : {location or 'n.c.'}\n"
            f"Description :\n{(jd_text or '(fiche non disponible : appuie-toi sur le titre et l’employeur)')[:7000]}"
            f"{extra}\n\nRédige le texte de candidature en {'français' if lang == 'FR' else 'anglais'} et rends le JSON.")
    lo, hi = WORDS[lang]
    letter = Letter.model_validate(
        claude.json(system=system, user=user, schema=LETTER_SCHEMA, effort="high", max_tokens=10000))
    n = word_count(letter.lettre)
    if n < lo or n > hi:
        log.info("lettre hors fourchette (%d mots) — nouvel essai", n)
        user2 = (user + f"\n\nTa version précédente faisait {n} mots : hors fourchette ({lo}-{hi}). "
                 "Réécris en respectant strictement la longueur, sans perdre les faits durs ni les clauses.")
        letter = Letter.model_validate(
            claude.json(system=system, user=user2, schema=LETTER_SCHEMA, effort="high", max_tokens=10000))
    if word_count(letter.lettre) < 120:
        raise RuntimeError(f"lettre trop courte ({word_count(letter.lettre)} mots) — dossier abandonné")
    return letter


def rewrite_letter(claude: Claude, *, lang: str, text: str, consigne: str, title: str, employer: str,
                   location: str, profile_md: str, writing_rules: str) -> Letter:
    """Réécrit une lettre existante selon une consigne de Xavier, mêmes garde-fous que la rédaction."""
    lang = "FR" if lang == "FR" else "EN"
    system = [prompt("cover_common"), prompt("cover_fr" if lang == "FR" else "cover_en"),
              profile_md, (writing_rules or "")]
    lo, hi = WORDS[lang]
    user = (f"OFFRE\nEmployeur : {employer}\nPoste : {title}\nLieu : {location or 'n.c.'}\n\n"
            f"LETTRE ACTUELLE :\n{text}\n\n"
            f"CONSIGNE DE XAVIER : {consigne.strip()}\n\n"
            f"Réécris la lettre en appliquant cette consigne, en {'français' if lang == 'FR' else 'anglais'}, "
            f"entre {lo} et {hi} mots, en conservant les faits durs verbatim et sans rien inventer. "
            "Ne change que ce que la consigne demande. Rends le JSON (lettre, objections).")
    letter = Letter.model_validate(
        claude.json(system=system, user=user, schema=LETTER_SCHEMA, effort="high", max_tokens=10000))
    if word_count(letter.lettre) < 100:
        raise RuntimeError(f"réécriture trop courte ({word_count(letter.lettre)} mots)")
    return letter


# ---------------------------------------------------------------------------
# DOCX au letterhead Xavier Advisory
# ---------------------------------------------------------------------------


def format_date(lang: str, day: dt.date | None = None) -> str:
    d = day or dt.date.today()
    if lang == "FR":
        mois = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
                "septembre", "octobre", "novembre", "décembre"]
        return f"{d.day} {mois[d.month - 1]} {d.year}"
    months = ["January", "February", "March", "April", "May", "June", "July", "August",
              "September", "October", "November", "December"]
    return f"{d.day} {months[d.month - 1]} {d.year}"


def _bottom_border(paragraph, color: str = "C79A3B", size: int = 12) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def _run(p, text: str, *, size: float, bold: bool = False, color: RGBColor | None = None, font: str | None = None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    if font:
        r.font.name = font
    return r


def strip_signature(text: str) -> str:
    """Retire la signature finale si le texte la porte déjà (la skill `cover-letter` signe la lettre, le DOCX aussi) :
    sinon le nom apparaîtrait deux fois."""
    lines = text.rstrip().split("\n")
    while lines and lines[-1].strip().lower() in (CONTACT_NAME.lower(), ""):
        lines.pop()
    return "\n".join(lines)


def letter_docx(text: str, out: Path, *, employer: str, title: str, lang: str, day: dt.date | None = None) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(56)
        s.left_margin = s.right_margin = Pt(62)
    normal = doc.styles["Normal"]
    normal.font.name = "Montserrat"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BRAND_NAVY
    normal.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    _run(p, CONTACT_NAME, size=22, bold=True, color=BRAND_NAVY, font="Playfair Display")
    p = doc.add_paragraph()
    _run(p, CONTACT_ROLE, size=8, color=BRAND_GOLD)
    p = doc.add_paragraph()
    _run(p, f"{CONTACT_EMAIL}   ·   {CONTACT_PHONE}", size=8.5, color=INK_SOFT)
    _bottom_border(p)
    p.paragraph_format.space_after = Pt(16)

    if employer:
        p = doc.add_paragraph()
        _run(p, employer, size=10.5, bold=True)
    p = doc.add_paragraph()
    _run(p, format_date(lang, day), size=9.5, color=INK_SOFT)
    p.paragraph_format.space_after = Pt(12)
    if title:
        p = doc.add_paragraph()
        _run(p, f"{'Objet' if lang == 'FR' else 'Re'} : {title}", size=10.5, bold=True)
        p.paragraph_format.space_after = Pt(12)

    text = strip_signature(text)
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text.strip()) if b.strip()]
    for b in blocks or [text.strip()]:
        p = doc.add_paragraph(re.sub(r"\s*\n\s*", " ", b))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.25

    # Signature : même typographie que le corps (correction Xavier du 29/07/2026)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.add_run(CONTACT_NAME)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def output_name(lang: str, employer: str, job_id: str, day: str) -> str:
    safe = re.sub(r"[^\w\-]", "", (employer or "offre").replace(" ", ""))[:30] or "offre"
    jid = f"_{job_id}" if job_id else ""
    prefix = "LM" if lang == "FR" else "CL"
    return f"{prefix}_XRO_{lang}_{safe}{jid}_{day}"
