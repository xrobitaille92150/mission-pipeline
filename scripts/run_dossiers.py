#!/usr/bin/env python3
"""
run_dossiers.py — Génère CV + Cover Letter pour les offres Veille en attente.
Tourne en headless via launchd à 6h, 12h, 19h.
Usage : python3 run_dossiers.py
"""

import os, sys, json, subprocess, datetime, re, time, logging, shutil, textwrap
import requests
from docx import Document
import anthropic

# ── Logging ────────────────────────────────────────────────────────────────────
LOG_DIR = os.path.expanduser("~/Desktop/Claude/Projects/Candidatures/pipeline/logs")
os.makedirs(LOG_DIR, exist_ok=True)
ts = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)s  %(message)s",
    handlers=[
        logging.FileHandler(f"{LOG_DIR}/dossiers_{ts}.log"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# ── Config ─────────────────────────────────────────────────────────────────────
def load_env():
    env_file = os.path.expanduser("~/.config/mission-pipeline/anthropic.env")
    if os.path.exists(env_file):
        for line in open(env_file):
            line = line.strip()
            if line and "=" in line and not line.startswith("#"):
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip())

load_env()

ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
AIRTABLE_PAT      = os.environ.get("AIRTABLE_PAT", "")

if not ANTHROPIC_API_KEY:
    log.error("ANTHROPIC_API_KEY manquant — arrêt.")
    sys.exit(1)
if not AIRTABLE_PAT:
    log.error("AIRTABLE_PAT manquant — arrêt.")
    sys.exit(1)

AIRTABLE_BASE   = "apphTpnW5vu0OdnfC"
AIRTABLE_TABLE  = "tblrCyL6huHkUPZbF"
GITHUB_REPO     = "xrobitaille92150/mission-pipeline"
REPO_PATH       = os.path.expanduser("~/Desktop/Claude/Projects/Candidatures/pipeline")
CV_BASE_DIR     = os.path.expanduser("~/Desktop/Claude/Projects/CV_Profiles/CV de base")
CV_OUT_DIR      = os.path.expanduser("~/Desktop/Claude/Projects/CV_Profiles")
PDF_DIR         = os.path.expanduser("~/Desktop/Claude/Projects/CV_Profiles/pdf")
# Copie éditable (.docx) vers le miroir Google Drive — filet si le compte GitHub
# (flaggé, ticket #4530517) saute, et base de retouches manuelles. Fail-soft.
DRIVE_DOSSIERS  = os.path.expanduser(
    "~/Mon Drive/XavierAdvisory/10_Work/Candidatures/dossiers")
CONTEXT_FILE    = os.path.expanduser("~/Desktop/Claude/Projects/_shared/context.md")
WRITING_EN      = os.path.expanduser("~/Desktop/Claude/Skills/WRITING RULES.md")
WRITING_FR      = os.path.expanduser("~/Desktop/Claude/Skills/REGLES-ECRITURE-FR.md")
CHROME          = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# ── Modèles Claude ───────────────────────────────────────────────────────────────
HAIKU_MODEL  = "claude-haiku-4-5-20251001"  # gap analysis CV (tâche mécanique)
SONNET_MODEL = "claude-sonnet-5"            # rédaction cover letter (tâche fine)

# ── Charte graphique (Brand Book Xavier Advisory) ────────────────────────────────
BRAND_NAVY    = "#0B1530"
BRAND_GOLD    = "#C79A3B"
BRAND_INK_SOFT = "rgba(11,21,48,0.62)"
CONTACT_NAME  = "Xavier Robitaille"
CONTACT_EMAIL = "xro@xavier-robitaille.com"
CONTACT_PHONE = "+33 6 64 89 09 43"

os.makedirs(PDF_DIR, exist_ok=True)

# Airtable field IDs
F = {
    "preparer":  "fldqC5lJmSG7lkiCO",
    "jepostule": "fld57QamoiPvafq2N",
    "cv":        "fldhRmOci5ofoY34i",
    "cl":        "fldXC5R5KHMLqwvbT",
    "ecarte":    "fld4Feuwx9SdWRqqO",
    "jobid":     "fldwMEnsfCi52BAU5",
    "employeur": "flddiP9RZb61Krm1u",
    "poste":     "fld4jROYUneaOcHBL",
    "lieu":      "fldGMSSdey77gM1i5",
    "url":       "fldCX37QooP7wPWVW",
    "score":     "fldYHTOWNIp0zc6UI",
    "note_role": "fld22OVfgmAFi3a1z",
    "note_crit": "fld9PrzOUhZ2I7cXQ",
}

CV_FILES = {
    "FinanceTransformation": "CV_XRO_EN_FinanceTransformation_v4.docx",
    "AssetManagement":       "CV_XRO_EN_AssetManagement_v4.docx",
    "IFRS17SolvencyII":      "CV_XRO_EN_IFRS17_SolvencyII_v4.docx",
}

# Versions FR par cluster (générées par make_fr_base_cvs.py). Sélectionnées quand
# l'annonce est en français ; fallback sur l'EN si le fichier FR est absent.
CV_FILES_FR = {
    "FinanceTransformation": "CV_XRO_FR_FinanceTransformation_v4.docx",
    "AssetManagement":       "CV_XRO_FR_AssetManagement_v4.docx",
    "IFRS17SolvencyII":      "CV_XRO_FR_IFRS17_SolvencyII_v4.docx",
}

# ── Airtable helpers ───────────────────────────────────────────────────────────
def at_headers():
    return {"Authorization": f"Bearer {AIRTABLE_PAT}", "Content-Type": "application/json"}

def fetch_pending():
    """Retourne les offres avec (Préparer dossier=true OU Je postule=true), CV vide, J'écarte=false.
    « Préparer dossier » est coché auto à Score>=SEUIL ; « Je postule » et « Préparer dossier »
    peuvent aussi être cochés MANUELLEMENT par Xavier — les deux déclenchent la génération CV/CL
    (rétroactif : toute la table est scannée à chaque run, cf. JACK horaire du 28/07/2026).
    Plafond optionnel par run via env MATT_MAX (0 ou absent = illimité)."""
    url = f"https://api.airtable.com/v0/{AIRTABLE_BASE}/{AIRTABLE_TABLE}"
    formula = (
        f"AND(OR({{{F['preparer']}}}=1,{{{F['jepostule']}}}=1),"
        f"{{{F['cv']}}}='',"
        f"NOT({{{F['ecarte']}}}=1))"
    )
    # Passer fields[] comme liste de tuples + retourner par ID de champ
    base_params = [
        ("filterByFormula", formula),
        ("sort[0][field]", F["score"]),
        ("sort[0][direction]", "desc"),
        ("returnFieldsByFieldId", "true"),
    ]
    for fid in F.values():
        base_params.append(("fields[]", fid))

    records = []
    offset = None
    while True:
        params = base_params.copy()
        if offset:
            params.append(("offset", offset))
        r = requests.get(url, headers=at_headers(), params=params, timeout=30)
        r.raise_for_status()
        data = r.json()
        records.extend(data.get("records", []))
        offset = data.get("offset")
        if not offset:
            break
    cap = int(os.environ.get("MATT_MAX", "0") or 0)
    if cap > 0 and len(records) > cap:
        records = records[:cap]  # tri Score desc déjà appliqué — on traite les meilleures d'abord
    return records

def update_record(record_id, cv_link, cl_link):
    url = f"https://api.airtable.com/v0/{AIRTABLE_BASE}/{AIRTABLE_TABLE}/{record_id}"
    payload = {"fields": {F["cv"]: cv_link, F["cl"]: cl_link}}
    r = requests.patch(url, headers=at_headers(), json=payload, timeout=30)
    r.raise_for_status()

# ── Claude API helper ──────────────────────────────────────────────────────────
client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

def claude(prompt, model=HAIKU_MODEL, max_tokens=4096):
    msg = client.messages.create(
        model=model,
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}],
    )
    # Sonnet 5 peut renvoyer des blocs "thinking" avant le texte → prendre les blocs texte.
    parts = [b.text for b in msg.content if getattr(b, "type", None) == "text"]
    return "".join(parts).strip()

# ── CV selection ───────────────────────────────────────────────────────────────
def select_cv(jd_text):
    """Choisit le profil CV selon le decision tree de cv_profiles.md."""
    jd = jd_text.lower()
    # Cluster B — Investissement/actif (AssetManagement CV)
    am_kw = ["oms","simcorp","clearwater","aladdin","bloomberg aim","murex","linedata","wealthsuite",
             "front-to-back","investment platform","investment accounting","investment reporting",
             "comptabilité des placements","comptabilité investissement","reporting investissement",
             "abor","ibor","custodian","securities migration","fund accounting",
             "middle office","back office"]
    # Cluster A — Assurance/passif (IFRS17SolvencyII CV)
    ifrs_kw = ["ifrs 17","ifrs17","ifrs 9","ifrs9","solvency ii","solvency 2","qrt","orsa",
               "actuarial","technical provisions","dry-run","paa","gmm","ecl","reserving","solvency",
               "comptabilité technique","provisions techniques","reporting prudentiel","pillar 3"]
    # FR back-office / interfaces comptables (AO CV)
    ao_kw = ["migration comptable","bascule","apurement","comptes d'attente","recettes (uat)",
             "cahier des charges","cdc","amoa back-office","interfaces comptables"]

    if any(k in jd for k in ifrs_kw):
        return "IFRS17SolvencyII"
    if any(k in jd for k in am_kw):
        return "AssetManagement"
    return "FinanceTransformation"

# ── Détection de langue ──────────────────────────────────────────────────────────
def detect_language(text):
    """Détecte FR vs EN à partir du texte de l'offre (heuristique mots fréquents).
    Remplace l'ancien lang=cv_type=='AO' (toujours EN car AO supprimé). La cover
    letter doit être rédigée dans la langue de l'annonce (règle du skill cover-letter)."""
    if not text:
        return "EN"
    t = " " + re.sub(r"\s+", " ", text.lower()) + " "
    fr_markers = [" le ", " la ", " les ", " des ", " une ", " un ", " et ", " pour ",
                  " avec ", " vous ", " nous ", " dans ", " sur ", " au ", " du ",
                  " en ", " est ", " que ", " qui ", " poste ", " entreprise ",
                  " expérience ", " compétences ", " missions ", " recherche ",
                  " sein ", " notre ", " vos ", " équipe "]
    en_markers = [" the ", " and ", " for ", " with ", " you ", " we ", " to ", " of ",
                  " in ", " on ", " is ", " are ", " will ", " role ", " team ",
                  " experience ", " skills ", " company ", " our ", " your ",
                  " within ", " as ", " a ", " an "]
    fr = sum(t.count(m) for m in fr_markers)
    en = sum(t.count(m) for m in en_markers)
    return "FR" if fr > en else "EN"

# ── Clauses conditionnelles (géographie / nature du poste) ───────────────────────
IR35_SENTENCE = "As a contractor who is a non-UK Tax resident, based overseas, IR35 does not apply to me."

def geo_role_clauses(lieu, jd_text, poste):
    """Retourne la liste des précisions à intégrer selon la géographie et la nature
    du poste. Détection déterministe : le pays vient d'abord du champ Lieu (fiable),
    puis du JD en repli. La nature (interne vs mission) vient du JD + intitulé."""
    loc = (lieu or "").lower().strip()
    if not loc:
        loc = (jd_text or "")[:400].lower()
    text_all = f"{poste or ''} {jd_text or ''}".lower()

    uk = (" uk" in f" {loc}") or any(k in loc for k in [
        "united kingdom", "royaume-uni", "angleterre", "england", "scotland",
        "écosse", "ecosse", "wales", "pays de galles", "london", "londres",
        "manchester", "edinburgh", "glasgow", "birmingham", "leeds", "bristol",
        "liverpool", "cambridge", "oxford"])
    ie = any(k in loc for k in ["ireland", "irlande", "dublin", "cork", "galway", "limerick"])
    benelux = any(k in loc for k in [
        "belgium", "belgique", "belgië", "brussels", "bruxelles", "antwerp",
        "anvers", "gent", "ghent", "netherlands", "pays-bas", "nederland",
        "amsterdam", "rotterdam", "hague", "haye", "utrecht", "eindhoven",
        "luxembourg", "benelux"])
    france = ("france" in loc) or any(k in loc for k in [
        "paris", "lyon", "marseille", "lille", "toulouse", "bordeaux", "nantes",
        "nice", "strasbourg", "rennes", "montpellier", "suresnes", "défense", "defense"])

    freelance_sig = any(k in text_all for k in [
        "freelance", "mission", "indépendant", "independant", "consultant externe",
        "prestation", "prestataire", "contractor", "contract role", "b2b",
        "daily rate", "tjm", "interim", "intérim", "sous-traitance"])

    hybride_clause = ("REMOTE/HYBRIDE : intègre naturellement que le travail hybride est une "
                      "pratique courante pour Xavier et qu'il maîtrise parfaitement les outils "
                      "et pratiques du travail à distance (référence concrète : sa dernière "
                      "expérience chez Clearwater).")

    clauses = []
    if uk:
        clauses.append(hybride_clause)
        clauses.append('IR35 : inclus la phrase EXACTE ci-dessous, en anglais, telle quelle — '
                       'ne la traduis pas, ne la reformule pas, ne la coupe pas : "' + IR35_SENTENCE + '"')
    elif ie or benelux:
        clauses.append(hybride_clause)

    if france and not freelance_sig:
        clauses.append("OUVERTURE CDI : ce poste est un rôle en interne en France. Précise, sans "
                       "lourdeur, que bien que Xavier travaille aujourd'hui en freelance, il reste "
                       "ouvert à un CDI — soit immédiatement, soit après une période initiale en mission.")
    return clauses

# ── LinkedIn fetch ─────────────────────────────────────────────────────────────
def fetch_linkedin(job_id):
    url = f"https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{job_id}"
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"}
    try:
        r = requests.get(url, headers=headers, timeout=20)
        if r.status_code != 200:
            return None
        # Extraire la description principale
        html = r.text
        m = re.search(r'class="show-more-less-html__markup[^"]*">(.*?)</div>', html, re.DOTALL)
        if m:
            raw = re.sub(r"<[^>]+>", " ", m.group(1))
            return re.sub(r"\s+", " ", raw).strip()[:6000]
        # Fallback : extraire tout le texte
        raw = re.sub(r"<[^>]+>", " ", html)
        return re.sub(r"\s+", " ", raw).strip()[:6000]
    except Exception as e:
        log.warning(f"LinkedIn fetch échoué ({e})")
        return None

# ── python-docx : retouches chirurgicales ─────────────────────────────────────
def apply_cv_edits(base_path, out_path, edits):
    """
    edits = liste de (old_text, new_text) à remplacer dans les paragraphes.
    Préserve la mise en forme. Ne modifie que les runs contenant old_text.
    """
    doc = Document(base_path)

    def replace_in_para(para, old, new):
        full = "".join(r.text for r in para.runs)
        if old in full:
            for r in para.runs:
                r.text = r.text.replace(old, new)
            return True
        return False

    for old, new in edits:
        if old == new:
            continue
        found = False
        for para in doc.paragraphs:
            if replace_in_para(para, old, new):
                found = True
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if replace_in_para(para, old, new):
                            found = True
        if not found:
            log.debug(f"Texte non trouvé pour remplacement : '{old[:40]}'")

    doc.save(out_path)

# ── PDF conversion ─────────────────────────────────────────────────────────────
CSS = """
@page { size: A4; margin: 11mm 13mm; }
body { font-size: 8.8pt; line-height: 1.26; font-family: "Helvetica Neue", Arial, sans-serif; }
h1, h2, h3 { color: #0b3d62; margin: 4pt 0 2pt; }
p, li { margin: 1pt 0; }
"""

def docx_to_pdf(docx_path, pdf_path):
    css_path = "/tmp/cv_dossiers_style.css"
    html_path = "/tmp/cv_dossiers_tmp.html"
    with open(css_path, "w") as f:
        f.write(CSS)
    subprocess.run([
        "/usr/local/bin/pandoc", docx_path, "-o", html_path,
        "--standalone", f"--css={css_path}",
        "--embed-resources", '--metadata', 'title=CV Xavier Robitaille'
    ], check=True, capture_output=True)
    subprocess.run([
        CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}", html_path
    ], check=True, capture_output=True)
    log.info(f"PDF généré : {pdf_path}")

def _format_date(lang):
    d = datetime.date.today()
    if lang == "FR":
        mois = ["janvier","février","mars","avril","mai","juin","juillet","août",
                "septembre","octobre","novembre","décembre"]
        return f"{d.day} {mois[d.month-1]} {d.year}"
    mois = ["January","February","March","April","May","June","July","August",
            "September","October","November","December"]
    return f"{mois[d.month-1]} {d.day}, {d.year}"

def _esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def text_to_pdf(text, pdf_path, employeur="", poste="", lang="EN",
                title="Cover Letter Xavier Robitaille"):
    """Génère un PDF de cover letter au letterhead Xavier Advisory (charte graphique)."""
    html_path = "/tmp/cl_dossiers_tmp.html"

    # Corps : découpe en paragraphes sur les lignes vides ; \n simples → espaces.
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text.strip()) if b.strip()]
    paras = [_esc(re.sub(r"\s*\n\s*", " ", b)) for b in blocks]
    body_html = "\n".join(f"<p>{p}</p>" for p in paras) or f"<p>{_esc(text)}</p>"

    date_str = _format_date(lang)
    subject_label = "Objet" if lang == "FR" else "Re"
    subject = f"{subject_label} : {_esc(poste)}" if poste else ""
    recipient = _esc(employeur)
    signoff = "Xavier Robitaille"

    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<style>
@import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Montserrat:wght@300;400;500;600&display=swap');
:root{{--navy:{BRAND_NAVY};--gold:{BRAND_GOLD};--ink-soft:{BRAND_INK_SOFT};}}
@page{{size:A4;margin:20mm 22mm}}
*{{box-sizing:border-box}}
body{{margin:0;color:var(--navy);
  font-family:'Montserrat',system-ui,'Helvetica Neue',Arial,sans-serif;
  font-size:10.5pt;line-height:1.6;-webkit-font-smoothing:antialiased}}
.letterhead{{display:flex;justify-content:space-between;align-items:flex-end;
  border-bottom:2px solid var(--gold);padding-bottom:10pt;margin-bottom:4pt}}
.name{{font-family:'Playfair Display',Georgia,serif;font-weight:700;
  font-size:22pt;letter-spacing:.5px;line-height:1}}
.role{{font-size:8pt;font-weight:500;letter-spacing:2.5px;text-transform:uppercase;
  color:var(--gold);margin-top:6pt}}
.contact{{text-align:right;font-size:8.5pt;color:var(--ink-soft);line-height:1.5}}
.contact .sep{{color:var(--gold)}}
.meta{{margin:20pt 0 2pt;font-size:9.5pt;color:var(--ink-soft)}}
.recipient{{font-weight:600;color:var(--navy)}}
.subject{{margin:14pt 0 16pt;font-weight:600;color:var(--navy);font-size:10.5pt}}
.body p{{margin:0 0 11pt;text-align:justify}}
.signoff{{margin-top:22pt}}
</style></head><body>
<div class="letterhead">
  <div>
    <div class="name">{CONTACT_NAME}</div>
    <div class="role">Finance &amp; Insurance Transformation</div>
  </div>
  <div class="contact">
    {CONTACT_EMAIL}<br>
    {CONTACT_PHONE}
  </div>
</div>
<div class="meta">
  {f'<span class="recipient">{recipient}</span><br>' if recipient else ''}{date_str}
</div>
{f'<div class="subject">{subject}</div>' if subject else '<div style="height:12pt"></div>'}
<div class="body">
{body_html}
</div>
<div class="signoff">{signoff}</div>
</body></html>"""
    with open(html_path, "w") as f:
        f.write(html)
    subprocess.run([
        CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}", html_path
    ], check=True, capture_output=True)
    log.info(f"PDF CL généré : {pdf_path}")

def cl_to_docx(text, docx_path, employeur="", poste="", lang="EN"):
    """Version .docx éditable de la cover letter (même contenu que le PDF).
    Mise en page sobre reprenant la charte : Playfair pour le nom, Montserrat
    pour le corps, filet or sous l'en-tête impossible en python-docx simple →
    on s'en tient au texte structuré (l'original de référence reste le PDF)."""
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    navy = RGBColor.from_string(BRAND_NAVY.lstrip("#"))
    gold = RGBColor.from_string(BRAND_GOLD.lstrip("#"))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Montserrat"
    style.font.size = Pt(10.5)
    style.font.color.rgb = navy

    p = doc.add_paragraph()
    r = p.add_run(CONTACT_NAME)
    r.font.name = "Playfair Display"
    r.font.size = Pt(22)
    r.bold = True
    p = doc.add_paragraph()
    r = p.add_run("FINANCE & INSURANCE TRANSFORMATION")
    r.font.size = Pt(8)
    r.font.color.rgb = gold
    p = doc.add_paragraph()
    r = p.add_run(f"{CONTACT_EMAIL}  ·  {CONTACT_PHONE}")
    r.font.size = Pt(8.5)

    doc.add_paragraph()
    if employeur:
        p = doc.add_paragraph()
        p.add_run(employeur).bold = True
    doc.add_paragraph(_format_date(lang))
    if poste:
        subject_label = "Objet" if lang == "FR" else "Re"
        p = doc.add_paragraph()
        p.add_run(f"{subject_label} : {poste}").bold = True
    doc.add_paragraph()

    blocks = [b.strip() for b in re.split(r"\n\s*\n", text.strip()) if b.strip()]
    for b in blocks or [text.strip()]:
        p = doc.add_paragraph(re.sub(r"\s*\n\s*", " ", b))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Signature : même typographie que le corps (correction Xavier 29/07)
    doc.add_paragraph()
    doc.add_paragraph(CONTACT_NAME)

    doc.save(docx_path)
    log.info(f"  CL .docx sauvegardée : {os.path.basename(docx_path)}")

def copy_to_drive(files):
    """Copie les fichiers vers le miroir Drive (10_Work/Candidatures/dossiers/
    YYYY-MM-DD/). Non bloquant : un échec ne fait pas échouer le dossier."""
    today = datetime.date.today().strftime("%Y-%m-%d")
    dest_dir = os.path.join(DRIVE_DOSSIERS, today)
    try:
        os.makedirs(dest_dir, exist_ok=True)
        for f in files:
            shutil.copy2(f, os.path.join(dest_dir, os.path.basename(f)))
        log.info(f"  Copie Drive OK : {len(files)} fichier(s) → dossiers/{today}/")
    except Exception as e:
        log.warning(f"  Copie Drive échouée (non bloquant) : {e}")

def build_cl_prompt(employeur, poste, lieu, description, context_md, writing_rules,
                    cl_lang, precisions_block=""):
    """Prompt CL v2.4 — standards distincts EN/FR + faits durs validés par Xavier (29/07/2026).
    Source : skill cover-letter v2.4 + skill voix-xavier (registre cover text, calibré v1)."""
    hard_facts = """═══ FAITS DURS — formulations VALIDÉES par Xavier, à utiliser verbatim, ne JAMAIS improviser ═══
1. TITRE D'ACTUAIRE INTERDIT (illégal en France — épreuves du CEA validées, mémoire non soutenu). Jamais « actuaire » / « qualification actuarielle » / "actuarial qualification". Autorisé — FR : « j'ai validé l'ensemble des épreuves du Centre d'Études Actuarielles (2004) » ; EN : "I completed the full examination track of the Centre d'Études Actuarielles (2004)".
2. CNP — Solvabilité II est prudentiel, PAS comptable. Jamais « comptabilité sous Solvabilité II ». FR : « Chez CNP, j'ai automatisé la production financière projetée multi-actifs (P&L, cash flows) et celle des états Solvabilité II relatifs aux placements : 350 Md€ d'encours, implémentation SimCorp Dimension de la conception à la mise en production, livrée en 18 mois. » EN : "At CNP, I automated the projected financial production across asset classes (P&L, cash flows) and the Solvency II investment reporting: €350bn in assets, SimCorp Dimension from design to go-live, delivered in 18 months."
3. Coface IFRS 9 — toujours « déployé par étapes sur quatre ans : POC successifs, méthodologie ECL, schémas comptables, recette, coordination avec le dépositaire CACEIS » / EN "rolled out in stages over four years". Data : « création chez Coface d'un Data Warehouse placements (champs, structure des tables, contrôle des inputs et des outputs) ».
4. Clearwater — « directeur de programme senior » / "senior programme director" ; « j'ai piloté en parallèle plusieurs implémentations d'une solution R to R en mode SaaS, avec des comités sponsors mensuels de niveau direction générale et une gestion consolidée des livrables, jalons et ressources sur un outil dédié (Monday.com) ». Client luxembourgeois : Luxempart.
5. SCOR — « quatre ans chez SCOR, dont deux comme chief transformation officer de la fonction investissement » (jamais « conseiller du COO »). Réassurance dès Mazars : SCOR, entités de réassurance AXA, entités en run-off.
6. Primexis — référence n°1 practice building / développement commercial : « création de la BU Assurance from scratch : 20 personnes, 10 clients, 3 M€ de CA, membre du Comex ».
7. Le MANAGEMENT n'est JAMAIS un écart (BU 20 personnes chez Primexis, équipes projet de 10 à 15, 16 collaborateurs chez Coface).
8. IA : « Certifié Generative AI Primer (Vanderbilt, 2026) et GenAI for Account Executives (Yale, 2026), j'automatise aujourd'hui la quasi-totalité de mon activité entrepreneuriale via des agents IA sous Claude ou GPT. »
9. Statut — FR : « J'interviens aujourd'hui en indépendant ; toutefois, un CDI est tout à fait envisageable pour moi, au regard de l'intérêt et des perspectives de ce poste. » EN : "I operate as an independent consultant today. Nevertheless, I am perfectly open to a permanent role, directly or after an initial engagement period."
10. ÉCARTS : ne JAMAIS auto-déclarer un écart technique (« je ne suis pas expert X ») ni logistique (localisation, mobilité, relocation). Seuls écarts admis : statut freelance/CDI (n°9) et calibrage de grade explicite dans l'annonce, via « Si [option A]… ; dans le cas contraire, je peux tout aussi bien [option B]. »
Invariants : « en assurance et en finance » (répéter la préposition) ; article devant les noms d'équipes (« l'équipe FAAS », jamais « FAAS » seul) ; « mobilisables chez vos clients » (jamais « opposables ») ; ne jamais énumérer « ORSA, RSR/SFCR, QRT » à côté de « Solvabilité II » (ces livrables SONT Solvabilité II → « les livrables Solvabilité II — ORSA, RSR/SFCR, QRT »)."""

    if cl_lang == "FR":
        regles_langue = """═══ STANDARD FRANÇAIS — LETTRE DE MOTIVATION (ne PAS calquer le standard anglo-saxon) ═══
LONGUEUR (contrainte DURE) : 190 à 280 mots. Ne descends JAMAIS sous 180.
STRUCTURE « Vous-Moi-Nous », registre formel soutenu, texte AÉRÉ (un paragraphe par idée ou preuve, séparés par une ligne vide) :
1. VOUS (2-3 phrases) — l'enjeu précis de l'entreprise, énoncé directement (exemple validé : « L'enjeu auquel doit faire face la direction financière de X est peu commun : … »). JAMAIS commencer par Xavier. Jamais d'abstraction contorsionnée ni de louange générique. Si l'entreprise est leader de son marché, le dire comme motivation.
2. MOI (l'essentiel) — 2 à 3 preuves chiffrées et développées (contexte, action, résultat), chacune reliée à un besoin de l'annonce, en utilisant les FAITS DURS verbatim. Un paragraphe par preuve, à la ligne entre chaque.
3. STATUT (si pertinent) — formulation validée du fait dur n°9.
4. NOUS + CLÔTURE, chaque élément SUR SA PROPRE LIGNE, dans cet ordre :
   - une phrase de projection (exemple validé : « La combinaison de mes diverses expériences me permettrait d'être opérationnel et force de proposition dès les premiers jours. ») ;
   - une phrase d'échange complète : « Je serais heureux d'avoir l'opportunité d'en discuter avec vous. » ou « Je suis à votre disposition pour échanger sur cette opportunité. » ou « Je suis disponible pour un entretien. » ;
   - ligne vide, puis « Bien cordialement. » seul sur sa ligne (OBLIGATOIRE — une clôture sans formule de politesse est cavalière en français, mais elle ne doit jamais terminer un paragraphe).
Pas de « Madame, Monsieur », pas de signature (le PDF s'en charge)."""
    else:
        regles_langue = """═══ ENGLISH STANDARD — COVER TEXT ═══
LENGTH (hard constraint): 250 to 350 words. Never under 220.
STRUCTURE, direct and confident:
1. COMPANY PROPOSITION (2-3 sentences) — the specific challenge from the JD. Never open with Xavier. No generic praise. If the company is the market leader, use it as motivation (validated example: "Swiss Re sets the standard in reinsurance, and that is the environment where I want to put this experience to work.").
2. EVIDENCE (the bulk) — 2-3 developed, quantified proof points tied to stated JD needs, using the HARD FACTS verbatim, plus one clause on why this company specifically. No vague referents ("winning new clients", never "the next ones"). Vary sentence openings — not every sentence starts with "I".
3. STATUS (if relevant) — validated formulation from hard fact 9.
4. CLOSE, each element ON ITS OWN LINE:
   - one full sentence, subject-verb-complement: "I would be happy to discuss [specific angle]." (banned: "Worth a conversation…?", "Happy to chat", any interrogative close — too familiar);
   - blank line, then "Best regards," alone on its line.
No salutation, no signature block (the PDF letterhead handles it)."""

    return f"""Tu rédiges un TEXTE DE CANDIDATURE pour Xavier Robitaille, consultant senior indépendant (25 ans d'expérience en assurance et en finance). Texte fluide, professionnel, destiné à être collé dans un formulaire de candidature ou envoyé comme corps de message.

{hard_facts}

{regles_langue}

═══ OFFRE ═══
Employeur : {employeur}
Poste : {poste}
Lieu : {lieu}
Description :
{description[:6000]}

═══ PROFIL DE XAVIER (source de vérité — ne rien inventer au-delà) ═══
{context_md[:4000]}

═══ RÈGLES D'ÉCRITURE (contraintes DURES, à appliquer intégralement) ═══
{writing_rules}

═══ INTERDITS ═══
- Commencer par Xavier / ses années d'expérience.
- Credential dump (lister tous les postes) ; paraphraser le CV.
- Louange d'entreprise générique ; référent vague.
- Auto-déclarer un écart technique ou logistique (fait dur n°10).
- Inventer un chiffre ou une expérience absents du profil ci-dessus.
- Titre, en-tête, « Madame, Monsieur », bloc signature.
{precisions_block}
Réponds UNIQUEMENT avec le texte de candidature, rien d'autre."""


# ── GitHub push ────────────────────────────────────────────────────────────────
def push_to_github(files, commit_msg):
    """files = liste de chemins locaux à copier dans candidatures/YYYY-MM-DD/."""
    today = datetime.date.today().strftime("%Y-%m-%d")
    dest_dir = os.path.join(REPO_PATH, "candidatures", today)
    os.makedirs(dest_dir, exist_ok=True)
    pushed = {}
    for f in files:
        name = os.path.basename(f)
        dst = os.path.join(dest_dir, name)
        shutil.copy2(f, dst)
        pushed[name] = f"https://github.com/{GITHUB_REPO}/blob/main/candidatures/{today}/{name}"
    subprocess.run(["git", "-C", REPO_PATH, "add", "candidatures/"], check=True, capture_output=True)
    subprocess.run(["git", "-C", REPO_PATH, "commit", "-m", commit_msg],
                   check=True, capture_output=True)
    subprocess.run(["git", "-C", REPO_PATH, "push", "origin", "main"],
                   check=True, capture_output=True)
    log.info(f"GitHub push OK : {list(pushed.values())}")
    return pushed

# ── Traitement d'une offre ─────────────────────────────────────────────────────
def process_offer(rec):
    c = rec["fields"]
    record_id = rec["id"]
    job_id    = c.get(F["jobid"], "")
    employeur = c.get(F["employeur"], "Inconnu")
    poste     = c.get(F["poste"], "")
    lieu      = c.get(F["lieu"], "")
    note_role = c.get(F["note_role"], "")
    note_crit = c.get(F["note_crit"], "")
    score = c.get(F["score"], "")

    log.info(f"→ {employeur} — {poste} [Score {score}]")
    today_str = datetime.date.today().strftime("%Y%m%d")
    safe_emp  = re.sub(r"[^\w\-]", "", employeur.replace(" ", ""))[:30]

    # 1. Fetch LinkedIn
    jd_text = None
    if job_id:
        jd_text = fetch_linkedin(job_id)
        if jd_text:
            log.info(f"  Description LinkedIn : {len(jd_text)} caractères")
        else:
            log.warning("  ⚠ Description LinkedIn non disponible — utilisation des notes")
        time.sleep(2)

    context_for_claude = jd_text or f"{note_role}\n\n{note_crit}" or f"{poste} chez {employeur} à {lieu}"

    # 2. Sélection CV + détection langue
    cv_type = select_cv(context_for_claude + " " + poste)
    # Langue de l'ANNONCE : détectée sur le JD (fiable), sinon l'intitulé, sinon les notes.
    # Elle pilote À LA FOIS la langue de la cover letter ET le choix du CV de base (FR/EN).
    cl_lang = detect_language(jd_text or poste or context_for_claude)
    log.info(f"  Langue annonce détectée : {cl_lang}")

    # CV FR si annonce FR et fichier FR présent ; sinon fallback EN.
    cv_lang = "EN"
    cv_base_file = CV_FILES[cv_type]
    if cl_lang == "FR":
        fr_file = CV_FILES_FR.get(cv_type)
        if fr_file and os.path.exists(os.path.join(CV_BASE_DIR, fr_file)):
            cv_lang = "FR"
            cv_base_file = fr_file
        else:
            log.warning(f"  CV FR absent pour {cv_type} — fallback CV EN")
    cv_base_path = os.path.join(CV_BASE_DIR, cv_base_file)
    if not os.path.exists(cv_base_path):
        log.error(f"  CV de base introuvable : {cv_base_path}")
        return None
    log.info(f"  CV sélectionné : {cv_type} ({cv_lang})")

    # 3. Gap analysis + retouches via Claude
    context_md = open(CONTEXT_FILE).read()[:8000] if os.path.exists(CONTEXT_FILE) else ""

    gap_prompt = f"""Tu aides à adapter un CV senior (25 ans d'expérience en assurance et finance) à une offre d'emploi.

OFFRE :
Employeur : {employeur}
Poste : {poste}
Lieu : {lieu}
Description : {context_for_claude[:3000]}

PROFIL :
{context_md[:3000]}

TÂCHE :
Identifie 3 à 6 retouches CHIRURGICALES à apporter au CV de base (type {cv_type}, rédigé en {"français" if cv_lang == "FR" else "anglais"}).
Chaque retouche = remplacer une courte phrase ou expression existante par une version légèrement améliorée.
Règles absolues : le texte "old" ET le texte "new" doivent être dans la MÊME langue que le CV ({"français" if cv_lang == "FR" else "anglais"}) ; ne pas inventer de chiffres ni d'expériences absentes du profil ; ne pas changer la structure.
Réponds UNIQUEMENT en JSON :
[
  {{"old": "texte exact à trouver dans le CV", "new": "texte de remplacement"}},
  ...
]
Si aucune retouche n'est utile, réponds : []
"""
    try:
        edits_raw = claude(gap_prompt, model=HAIKU_MODEL)
        # Extraire uniquement le bloc JSON [ ... ] pour éviter les textes parasites
        m = re.search(r'\[[\s\S]*\]', edits_raw)
        if not m:
            raise ValueError("Pas de JSON array dans la réponse")
        edits = [(e["old"], e["new"]) for e in json.loads(m.group()) if e.get("old") and e.get("new")]
        log.info(f"  {len(edits)} retouche(s) CV")
    except Exception as e:
        log.warning(f"  Gap analysis échoué ({e}) — CV sans retouche")
        edits = []

    # 4. Appliquer les retouches python-docx
    # Suffixe jobId : deux offres du même employeur le même jour ne s'écrasent plus (fix 29/07)
    jid = f"_{job_id}" if job_id else ""
    cv_out_name = f"CV_XRO_{cv_lang}_{cv_type}_{safe_emp}{jid}_{today_str}.docx"
    cv_out_path = os.path.join(CV_OUT_DIR, cv_out_name)
    apply_cv_edits(cv_base_path, cv_out_path, edits)
    log.info(f"  CV .docx sauvegardé : {cv_out_name}")

    # 5. Cover letter via Claude (Sonnet — rédaction fine, logique du skill cover-letter)
    writing_path = WRITING_FR if cl_lang == "FR" else WRITING_EN
    # Règles d'écriture chargées EN ENTIER (contrainte de style dure, pas d'aperçu tronqué).
    writing_rules = open(writing_path).read() if os.path.exists(writing_path) else ""

    # Clauses conditionnelles selon géographie / nature du poste
    clauses = geo_role_clauses(lieu, jd_text, poste)
    if clauses:
        precisions_block = ("\n═══ PRÉCISIONS À INTÉGRER (OBLIGATOIRE — selon la géographie/nature du poste) ═══\n"
                            + "\n".join(f"- {c}" for c in clauses)
                            + "\nIntègre ces précisions de façon FLUIDE et NATURELLE dans le corps ou la clôture — "
                              "jamais sous forme de liste, de puces ou de bloc à part. Elles s'ajoutent au texte : "
                              "tu peux dépasser la fourchette de ~40 mots pour les intégrer proprement.\n")
        log.info(f"  Clauses conditionnelles : {len(clauses)}")
    else:
        precisions_block = ""

    cl_prompt = build_cl_prompt(employeur, poste, lieu, context_for_claude,
                                context_md, writing_rules, cl_lang, precisions_block)
    # max_tokens élevé : Sonnet dépense une grande part du budget en blocs
    # "thinking" AVANT le texte — à 6000 le premier appel renvoyait 0 caractère
    # SYSTÉMATIQUEMENT (constat audit 28/07 : retry payant sur chaque CL).
    # 12000 direct = un seul appel ; les tokens non générés ne coûtent rien.
    cl_text = claude(cl_prompt, model=SONNET_MODEL, max_tokens=12000)
    if len(cl_text) < 400:
        log.warning(f"  Cover letter trop courte ({len(cl_text)} c.) — retry avec budget élargi")
        cl_text = claude(cl_prompt, model=SONNET_MODEL, max_tokens=16000)
    if len(cl_text) < 400:
        # Fail-loud : mieux vaut PAS de dossier qu'une CL vide poussée sur GitHub
        log.error(f"  Cover letter VIDE après retry ({len(cl_text)} c.) — dossier abandonné")
        return None
    log.info(f"  Cover letter générée ({cl_lang}, Sonnet) : {len(cl_text)} caractères")

    # 6. Conversion PDF
    cv_pdf_name = cv_out_name.replace(".docx", ".pdf")
    cv_pdf_path = os.path.join(PDF_DIR, cv_pdf_name)
    cl_pdf_name = f"CL_XRO_{cl_lang}_{safe_emp}{jid}_{today_str}.pdf"
    cl_pdf_path = os.path.join(PDF_DIR, cl_pdf_name)

    try:
        docx_to_pdf(cv_out_path, cv_pdf_path)
        text_to_pdf(cl_text, cl_pdf_path, employeur=employeur, poste=poste, lang=cl_lang)
    except subprocess.CalledProcessError as e:
        log.error(f"  Conversion PDF échouée : {e}")
        return None

    # 7. Push GitHub
    try:
        links = push_to_github(
            [cv_pdf_path, cl_pdf_path],
            f"dossiers: {employeur} {today_str}"
        )
        cv_link = links.get(cv_pdf_name, "")
        cl_link = links.get(cl_pdf_name, "")
    except subprocess.CalledProcessError as e:
        stderr = (e.stderr or b"").decode("utf-8", "replace").strip()
        stdout = (e.stdout or b"").decode("utf-8", "replace").strip()
        log.error(f"  Git push échoué : {e}\n  stderr: {stderr}\n  stdout: {stdout}")
        return None

    # 8. Mise à jour Airtable
    try:
        update_record(record_id, cv_link, cl_link)
        log.info(f"  ✓ Airtable mis à jour")
    except Exception as e:
        log.error(f"  Airtable update échoué : {e}")
        return None

    # 9. Copie .docx éditables vers le miroir Drive (fail-soft)
    cl_docx_path = os.path.join(CV_OUT_DIR,
                                f"CL_XRO_{cl_lang}_{safe_emp}{jid}_{today_str}.docx")
    try:
        cl_to_docx(cl_text, cl_docx_path, employeur=employeur, poste=poste,
                   lang=cl_lang)
    except Exception as e:
        log.warning(f"  CL .docx échouée (non bloquant) : {e}")
        cl_docx_path = None
    copy_to_drive([p for p in (cv_out_path, cl_docx_path) if p])

    return {"employeur": employeur, "cv": cv_link, "cl": cl_link}

# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    log.info("=== run_dossiers.py démarré ===")

    try:
        records = fetch_pending()
    except Exception as e:
        log.error(f"Airtable fetch échoué : {e}")
        sys.exit(1)

    if not records:
        log.info("Aucun dossier en attente. Fin.")
        return

    log.info(f"{len(records)} dossier(s) à générer")
    for rec in records:
        c = rec.get("fields", {})
        emp = c.get(F["employeur"], "?")
        pos = c.get(F["poste"], "?")
        log.info(f"  • {emp} — {pos}")

    results = []
    errors  = []
    for rec in records:
        try:
            res = process_offer(rec)
            if res:
                results.append(res)
            else:
                errors.append(rec["fields"].get(F["employeur"], "?"))
        except Exception as e:
            emp = rec["fields"].get(F["employeur"], "?")
            log.error(f"Erreur sur {emp} : {e}", exc_info=True)
            errors.append(emp)
        time.sleep(3)

    log.info(f"=== Terminé : {len(results)}/{len(records)} OK, {len(errors)} erreur(s) ===")
    for r in results:
        log.info(f"  ✓ {r['employeur']} | CV → {r['cv']}")
    for e in errors:
        log.info(f"  ✗ {e}")

if __name__ == "__main__":
    main()
